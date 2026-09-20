from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from lume_ingestion.pipeline import run_pipeline


def test_itau_current_account_xlsx_uses_printed_values_and_period_year(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "extrato"
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append([None, "Nome:", None, "Cliente Teste", "Agência/Conta:", "0191/ 41707-0"])
    sheet.append([])
    sheet.append([None, "Data:", None, "02/02/2026"])
    sheet.append([])
    sheet.append([None, "Extrato de Conta Corrente", None, None, None, "01/01/2026 a 02/02/2026"])
    sheet.append([])
    sheet.append([None, "Data", None, None, "Lançamento", "Valor (R$)", "Saldo (R$)"])
    sheet.append([None, "31/12", None, None, "SALDO ANTERIOR", None, 1000.00])
    sheet.append([None, "02/01", None, None, "PIX ENVIADO FORNECEDOR", -100.00, None])
    sheet.append([None, "02/01", None, None, "PIX TRANSF ITAU CLIENTE", 50.00, None])
    sheet.append([None, "02/01", None, None, "S A L D O", None, 950.00])
    path = tmp_path / "itau-conta-corrente.xlsx"
    workbook.save(path)

    result = run_pipeline(path, tmp_path / "output")

    assert result.success, result.errors
    assert result.source_format == "xlsx"
    assert result.document_type == "bank_statement"
    assert result.data["bank"] == "Itaú"
    assert result.data["branch"] == "0191"
    assert result.data["account"] == "41707-0"
    assert result.data["period_start"] == "2026-01-01"
    assert result.data["period_end"] == "2026-02-02"
    assert result.data["initial_balance"] == "1000.00"
    assert result.data["final_balance"] == "950.00"
    assert [row["date"] for row in result.data["transactions"]] == ["2026-01-02", "2026-01-02"]
    assert [row["amount"] for row in result.data["transactions"]] == ["-100.00", "50.00"]
    assert result.data["transactions"][0]["counterparty"] is None
    assert result.data["daily_balances"][0]["date"] == "2026-01-02"
    assert result.data["daily_balances"][0]["balance"] == "950.00"


def test_itau_spreadsheet_with_counterparty_columns_still_reads_razao_social(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Agência", "0191", "Conta", "0041707-0", "Itaú", "Período", "01/08/2026 a 25/08/2026"])
    sheet.append(["Data", "Lançamento", "Razão Social", "CPF/CNPJ", "Valor", "Saldo"])
    sheet.append(["01/08/2026", "SALDO ANTERIOR", None, None, None, 100.00])
    sheet.append(["01/08/2026", "PIX RECEBIDO", "CLIENTE LTDA", "00000000000191", 50.00, None])
    sheet.append(["01/08/2026", "SALDO EM CONTA CORRENTE", None, None, None, 150.00])
    path = tmp_path / "itau-razao.xlsx"
    workbook.save(path)

    result = run_pipeline(path, tmp_path / "output")

    assert result.success, result.errors
    assert result.data["transactions"][0]["counterparty"] == "CLIENTE LTDA"
    assert result.data["transactions"][0]["counterparty_tax_id"] == "00000000000191"
    assert result.data["branch"] == "0191"
    assert result.data["account"] == "0041707-0"


def test_itau_current_account_xlsx_keeps_saldo_anterior_across_a_weekend(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([None, "Agência/Conta:", "0191/ 41707-0"])
    sheet.append([None, "Extrato de Conta Corrente", "05/01/2026 a 09/01/2026"])
    sheet.append([None, "Data", "Lançamento", "Valor (R$)", "Saldo (R$)"])
    sheet.append([None, "02/01", "SALDO ANTERIOR", None, 200.00])
    sheet.append([None, "05/01", "PIX ENVIADO FORNECEDOR", -20.00, None])
    sheet.append([None, "05/01", "S A L D O", None, 180.00])
    path = tmp_path / "itau-weekend.xlsx"
    workbook.save(path)

    result = run_pipeline(path, tmp_path / "output")

    assert result.success, result.errors
    assert result.data["initial_balance"] == "200.00"
    assert result.data["transactions"][0]["date"] == "2026-01-05"


def test_itau_current_account_xlsx_without_period_does_not_succeed_empty(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([None, "Agência/Conta:", "0191/ 41707-0"])
    sheet.append([None, "Extrato de Conta Corrente"])
    sheet.append([None, "Data", "Lançamento", "Valor (R$)", "Saldo (R$)"])
    sheet.append([None, "02/01", "PIX ENVIADO FORNECEDOR", -20.00, None])
    path = tmp_path / "itau-sem-periodo.xlsx"
    workbook.save(path)

    result = run_pipeline(path, tmp_path / "output")

    assert not result.success
    assert result.errors[0].code == "unsupported_bank_statement_layout"


def test_xlsx_without_bank_or_cash_headers_still_fails_as_cash(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.append(["something else", "value"])
    path = tmp_path / "without-header.xlsx"
    workbook.save(path)

    result = run_pipeline(path, tmp_path / "output")

    assert not result.success
    assert result.errors[0].code == "spreadsheet_columns_not_mapped"
