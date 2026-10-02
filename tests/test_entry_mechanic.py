"""Generic entry: any readable file comes in; only the missing piece is asked.

Kinds stay apart.  A saved recipe applies only when that same kind still
matches.  A shifted column pauses.  An unread scan asks for the text,
instead of dying as an unknown type.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from pypdf import PdfWriter

from lume_ingestion.layouts import recipe_for
from lume_ingestion.layouts.cash_ledger_report import recipe_matches
from lume_ingestion.pipeline import run_pipeline

CASH_RECIPE = {
    "id": "cash-ledger-report-v1",
    "family": "cash_ledger",
    "bank": None,
    "source_formats": ["pdf", "xlsx"],
    "headers": ["DATA", "ENTRADA", "SAIDA", "SALDO"],
    "required": ["date", "inflow", "outflow", "balance"],
}

_HARD_STOPS = {
    "unknown_format",
    "unsupported_format",
    "unsupported_raw_format",
    "unsupported_normalized_format",
    "unexpected_processing_error",
    "cash_ledger_header_not_found",
    "page_without_useful_text",
}


def _pdf(path: Path, tokens: list[tuple[float, float, str]]) -> None:
    lines = ["BT", "/F1 8 Tf"]
    for x, top, text in tokens:
        y = 595 - top
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        lines.append(f"1 0 0 1 {x:.2f} {y:.2f} Tm ({escaped}) Tj")
    lines.append("ET")
    stream = ("\n".join(lines) + "\n").encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 842 595] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)


def _only(recipes):
    def _load():
        return recipes
    return _load


def _pause(result):
    assert not result.success
    assert result.errors, result.errors
    assert result.errors[0].code == "spreadsheet_columns_not_mapped"
    assert result.errors[0].code not in _HARD_STOPS
    assert "csv ou xlsx" not in result.errors[0].message.casefold()
    layout = result.errors[0].details["layout"]
    assert isinstance(layout, dict)
    return layout


def test_unknown_table_asks_only_the_unmapped_required_field(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("LUME_FORMAT_REGISTRY", str(tmp_path / "learned.json"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Data", "Historico", "Foo"])
    sheet.append(["13/02/2026", "PIX UNIVERSO", "nota"])
    path = tmp_path / "sem-valor.xlsx"
    workbook.save(path)

    result = run_pipeline(path, tmp_path / "out")

    proposal = _pause(result)
    assert proposal["missing"] == ["amount"]
    assert set(proposal["suggestions"].values()) >= {"date", "description"}
    assert "amount" not in proposal["suggestions"].values()
    assert "kind" not in proposal["missing"]
    assert "entrada" not in result.errors[0].message.casefold()
    assert "saldo" not in result.errors[0].message.casefold()
    assert "ledgers" not in result.data
    assert "transactions" not in result.data


def test_unknown_pdf_table_enters_and_asks_only_the_kind(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("LUME_FORMAT_REGISTRY", str(tmp_path / "learned.json"))
    path = tmp_path / "sem-rotulo.pdf"
    _pdf(path, [
        (40, 80, "Alpha"),
        (180, 80, "Beta"),
        (320, 80, "Gamma"),
        (40, 100, "xx"),
        (180, 100, "yy"),
        (320, 100, "zz"),
    ])

    result = run_pipeline(path, tmp_path / "out")

    proposal = _pause(result)
    assert proposal["missing"] == ["kind"]
    assert result.errors[0].message == "Que tipo de arquivo é este?"
    assert "ledgers" not in result.data


def test_odd_extension_still_enters_when_the_bytes_are_a_table(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("LUME_FORMAT_REGISTRY", str(tmp_path / "learned.json"))
    path = tmp_path / "movimentos.txt"
    path.write_text("Alpha;Beta\nxx;yy\n", encoding="utf-8")

    result = run_pipeline(path, tmp_path / "out")

    proposal = _pause(result)
    assert result.source_format == "csv"
    assert proposal["missing"] == ["kind"]
    assert result.errors[0].code != "unknown_format"


def test_chart_pdf_is_not_a_dead_rejection_and_ignores_the_cash_recipe(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("lume_ingestion.format_registry.load_formats", _only([CASH_RECIPE]))
    path = tmp_path / "plano.pdf"
    _pdf(path, [
        (40, 80, "Conta"),
        (180, 80, "Descricao"),
        (40, 100, "1.1.01"),
        (180, 100, "Caixa"),
    ])

    result = run_pipeline(path, tmp_path / "out")

    proposal = _pause(result)
    assert proposal["family"] == "chart"
    assert proposal["missing"] == []
    assert "code" in proposal["suggestions"].values()
    assert "description" in proposal["suggestions"].values()
    assert result.document_type != "cash_ledger"
    assert "ledgers" not in result.data
    assert recipe_for(
        {"source_format": "pdf", "document_type": "chart_of_accounts", "pages": [], "workbook": {"sheets": []}},
        [CASH_RECIPE],
    ) is None
    assert recipe_matches(
        {"source_format": "pdf", "document_type": "chart_of_accounts", "pages": []},
        [CASH_RECIPE],
    ) is None


def test_history_pdf_enters_as_history_not_as_cash(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("lume_ingestion.format_registry.load_formats", _only([CASH_RECIPE]))
    path = tmp_path / "historico.pdf"
    _pdf(path, [
        (30, 80, "Data"),
        (120, 80, "Debito"),
        (220, 80, "Credito"),
        (340, 80, "Valor"),
        (460, 80, "Complemento"),
        (30, 100, "02/01/2026"),
        (120, 100, "1.1.01"),
        (220, 100, "2.1.01"),
        (340, 100, "10,00"),
        (460, 100, "nota"),
    ])

    result = run_pipeline(path, tmp_path / "out")

    proposal = _pause(result)
    assert proposal["family"] == "history"
    assert result.document_type != "cash_ledger"
    assert "ledgers" not in result.data


def test_chart_spreadsheet_does_not_reuse_a_cash_recipe(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("lume_ingestion.format_registry.load_formats", _only([CASH_RECIPE]))
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Conta", "Descricao"])
    sheet.append(["1.1.01", "Caixa"])
    path = tmp_path / "plano.xlsx"
    workbook.save(path)

    result = run_pipeline(path, tmp_path / "out")

    assert result.success, result.errors
    assert result.document_type == "chart_of_accounts"
    assert result.data["accounts"][0]["code"] == "1.1.01"


def test_shifted_pinned_column_pauses_instead_of_reading_the_decoy(tmp_path: Path, monkeypatch) -> None:
    recipe = {
        "id": "cash-ledger-spreadsheet-v1",
        "family": "cash_ledger",
        "bank": None,
        "source_formats": ["xlsx"],
        "headers": ["DATA", "ENTRADA", "SAIDA", "SALDO"],
        "required": ["date", "inflow", "outflow", "balance"],
        "columns": {"date": 1, "inflow": 2, "outflow": 3, "balance": 4},
    }
    monkeypatch.setattr("lume_ingestion.format_registry.load_formats", _only([recipe]))
    workbook = Workbook()
    sheet = workbook.active
    # Saldo moved to column 5.  Column 4 is a decoy the old index would trust.
    sheet.append(["Data", "Entrada", "Saida", None, "Saldo"])
    sheet.append(["2026-01-02", 10, 0, 999, 10])
    path = tmp_path / "caixa-deslocado.xlsx"
    workbook.save(path)

    result = run_pipeline(path, tmp_path / "out")

    proposal = _pause(result)
    assert "balance" in proposal["shifted"]
    assert "balance" in proposal["suggestions"].values()
    assert "ledgers" not in result.data


def test_unreadable_scan_asks_for_the_text_instead_of_crashing(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("lume_ingestion.parsers.pdf_recover.recover_page", lambda *args, **kwargs: None)
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as stream:
        writer.write(stream)

    result = run_pipeline(path, tmp_path / "out")

    proposal = _pause(result)
    assert proposal["gap"] == "unreadable_text"
    assert proposal["missing"] == ["text"]
    assert "não consigo ler o texto" in result.errors[0].message.casefold()
    assert result.errors[0].code != "unexpected_processing_error"
