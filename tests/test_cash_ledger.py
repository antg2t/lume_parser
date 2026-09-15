from __future__ import annotations

import hashlib
import json
import copy
import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from lume_ingestion.artifacts import read_json
from lume_ingestion.cash_ledger import compare_cash_ledger_collections, normalize_cash_ledger_pdf
from lume_ingestion.pipeline import extract, run_pipeline, validate


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS_ROOT = PROJECT_ROOT / "docs"
XLSX = next(DOCS_ROOT.glob("Controle de caixa/xlsx/*.xlsx"))
PDF = next(DOCS_ROOT.glob("Controle de caixa/pdf/*.pdf"))
GOLDEN = json.loads((Path(__file__).parent / "goldens" / "cash_ledger.json").read_text(encoding="utf-8"))
ENTRY_FIELDS = ("date", "issue_date", "document", "counterparty", "notes", "inflow", "outflow", "balance")


def all_entries(data: dict) -> list[dict]:
    return [entry for ledger in data["ledgers"] for entry in ledger["entries"]]


@pytest.mark.integration
def test_xlsx_cash_ledger_matches_the_full_golden(tmp_path: Path) -> None:
    result = run_pipeline(XLSX, tmp_path / "output")

    assert result.success, result.errors
    assert result.source_format == "xlsx"
    assert result.document_type == "cash_ledger"
    entries = all_entries(result.data)
    payload = [{field: entry[field] for field in ENTRY_FIELDS} for entry in entries]
    digest = hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    assert len(entries) == GOLDEN["entry_count"]
    assert [len(ledger["entries"]) for ledger in result.data["ledgers"]] == GOLDEN["ledger_entry_counts"]
    assert digest == GOLDEN["entries_sha256"]
    assert payload[0] == GOLDEN["first_entry"]
    assert payload[-1] == GOLDEN["last_entry"]
    assert entries[0]["origin"]["sheet"] == "default.Table2"
    assert entries[0]["origin"]["cell_refs"] == ["A2", "B2", "D2", "E2", "F2", "G2", "H2"]
    assert any(entry["document"] is None for entry in entries)


@pytest.mark.integration
def test_pdf_and_xlsx_cash_ledgers_compare_without_divergences(tmp_path: Path) -> None:
    xlsx = run_pipeline(XLSX, tmp_path / "output")
    pdf = run_pipeline(PDF, tmp_path / "output")

    assert xlsx.success, xlsx.errors
    assert pdf.success, pdf.errors
    assert pdf.document_type == "cash_ledger"
    assert sum(len(ledger["entries"]) for ledger in pdf.data["ledgers"]) == GOLDEN["entry_count"]
    assert [ledger["account"] for ledger in pdf.data["ledgers"]] == [
        "237 -BRADESCO",
        "341 -ITAU",
        "341.3 -ITAU CONTA INVESTIMENTO",
        "N/D -N/D",
    ]
    assert all(entry["origin"]["page_number"] for entry in all_entries(pdf.data))
    assert any(entry["notes"] and "passivo" in entry["notes"] for entry in all_entries(pdf.data))
    assert "SECRETARIA" in next(entry for entry in all_entries(pdf.data) if entry["document"] == "ICMS-06/2026")["counterparty"]
    assert compare_cash_ledger_collections(xlsx.data, pdf.data) == {
        "matches": True,
        "reference_entry_count": 261,
        "candidate_entry_count": 261,
        "differences": [],
    }


@pytest.mark.integration
def test_pdf_page_with_only_a_repeated_header_is_reported_as_an_unexpected_break(tmp_path: Path) -> None:
    result = run_pipeline(PDF, tmp_path / "output")
    raw = copy.deepcopy(read_json(result.outputs.raw_json))
    raw["pages"][1]["words"] = [word for word in raw["pages"][1]["words"] if float(word["top"]) < 100]

    _, warnings = normalize_cash_ledger_pdf(raw)

    assert any(
        warning.code == "unexpected_cash_ledger_page_break" and warning.details["page_number"] == 2
        for warning in warnings
    )


@pytest.mark.integration
def test_cash_ledger_recognition_does_not_depend_on_vanguarda_or_a_bank_name(tmp_path: Path) -> None:
    renamed_xlsx = tmp_path / "any-ledger.xlsx"
    renamed_pdf = tmp_path / "another-report.pdf"
    shutil.copyfile(XLSX, renamed_xlsx)
    shutil.copyfile(PDF, renamed_pdf)

    xlsx = run_pipeline(renamed_xlsx, tmp_path / "output")
    pdf = run_pipeline(renamed_pdf, tmp_path / "output")

    assert xlsx.success and xlsx.document_type == "cash_ledger"
    assert pdf.success and pdf.document_type == "cash_ledger"
    assert [ledger["account"] for ledger in pdf.data["ledgers"]][-1] == "N/D -N/D"


def test_xlsx_without_cash_header_fails_with_a_precise_code(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["something else", "value"])
    path = tmp_path / "without-header.xlsx"
    workbook.save(path)

    result = run_pipeline(path, tmp_path / "output")

    assert not result.success
    assert result.errors[0].code == "cash_ledger_header_not_found"


def test_corrupted_xlsx_returns_a_structured_error(tmp_path: Path) -> None:
    workbook = Workbook()
    path = tmp_path / "corrupted.xlsx"
    workbook.save(path)
    with pytest.warns(UserWarning, match="Duplicate name"):
        with zipfile.ZipFile(path, "a") as package:
            package.writestr("xl/workbook.xml", "<workbook")

    result = run_pipeline(path, tmp_path / "output")

    assert not result.success
    assert result.errors[0].code == "invalid_xlsx"


def test_formula_without_cache_is_preserved_and_warned(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["DATA", "EMISSAO", "DOC.", "CLIENTE / FORNECEDOR", "ANOTACOES", "ENTRADA", "SAIDA", "SALDO"])
    sheet.append(["2026-08-01", "2026-08-01", None, "Fornecedor", None, 10, 0, "=F2-G2"])
    path = tmp_path / "formula.xlsx"
    workbook.save(path)

    raw, _ = extract(path, tmp_path / "output")

    assert raw["workbook"]["formula_count"] == 1
    assert any(warning["code"] == "formula_cache_missing" for warning in raw["warnings"])
    balance = raw["workbook"]["sheets"][0]["rows"][1]["cells"][-1]
    assert balance["formula"] == "=F2-G2"
    assert balance["calculated_value"] is None


def test_balance_divergence_is_an_error_with_its_spreadsheet_origin(tmp_path: Path) -> None:
    result = run_pipeline(XLSX, tmp_path / "output")
    normalized = read_json(result.outputs.normalized_json)
    normalized["cash_ledger"]["ledgers"][0]["entries"][1]["balance"] = "1.00"
    changed = tmp_path / "normalized.json"
    changed.write_text(json.dumps(normalized, ensure_ascii=False), encoding="utf-8")

    validation = validate(changed)

    assert not validation.success
    mismatch = next(error for error in validation.errors if error.code == "cash_balance_mismatch")
    assert mismatch.details["entry_index"] == 2
    assert mismatch.details["origin"]["row_number"] == 3
