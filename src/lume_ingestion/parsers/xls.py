"""RAW extraction for legacy XLS workbooks, kept read-only via xlrd."""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from time import perf_counter
from typing import Any

import xlrd
from openpyxl.utils import get_column_letter

from lume_ingestion.format_registry import layout_proposal_from_extract, match_extract
from lume_ingestion.errors import IngestionFailure
from lume_ingestion.models import SourceFile, Warning


def _raw_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


class XlsParser:
    name = "xls-xlrd"
    version = "0.1.0"

    def extract(
        self,
        path: Path,
        source: SourceFile,
        warnings: list[Warning],
        limits: Any = None,
    ) -> dict[str, Any]:
        del limits
        started = perf_counter()
        try:
            workbook = xlrd.open_workbook(path, on_demand=True)
        except (OSError, xlrd.XLRDError) as exc:
            raise IngestionFailure("invalid_xls", "A planilha XLS esta corrompida ou nao possui estrutura legivel.", reason=str(exc)) from exc
        sheets: list[dict[str, Any]] = []
        try:
            for sheet in workbook.sheets():
                rows: list[dict[str, Any]] = []
                for row_index in range(sheet.nrows):
                    cells: list[dict[str, Any]] = []
                    for column_index in range(sheet.ncols):
                        cell = sheet.cell(row_index, column_index)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode)
                        if value in (None, ""):
                            continue
                        column = column_index + 1
                        cells.append({
                            "coordinate": f"{get_column_letter(column)}{row_index + 1}",
                            "column": column,
                            "value": _raw_value(value),
                            "formula": None,
                            "calculated_value": None,
                        })
                    if cells:
                        rows.append({"row_number": row_index + 1, "cells": cells})
                sheets.append({
                    "name": sheet.name,
                    "dimensions": {"max_row": sheet.nrows, "max_column": sheet.ncols},
                    "merged_cells": [],
                    "rows": rows,
                })
        finally:
            workbook.release_resources()
        preview = {"source": source.model_dump(mode="json"), "source_format": "xls", "workbook": {"sheets": sheets}}
        matched = match_extract(preview, source.name)
        raw = {
            "schema_version": "1.0",
            "source": source.model_dump(mode="json"),
            "source_format": "xls",
            "document_type": matched.family if matched else None,
            "parser": self.name,
            "parser_version": self.version,
            "extraction_duration_ms": round((perf_counter() - started) * 1000),
            "workbook": {"sheet_count": len(sheets), "sheets": sheets, "formula_count": 0},
            "warnings": [warning.model_dump(mode="json") for warning in warnings],
            "errors": [],
        }
        if matched:
            raw["document_recognition"] = matched.as_recognition()
        else:
            raw["layout_proposal"] = layout_proposal_from_extract(raw)
        return raw
