"""RAW extraction for XLSX workbooks.

The workbook is opened read-only twice: once with formulas visible and once
with Excel's cached results visible.  This is intentionally kept separate
from cash-ledger normalization, so the RAW remains useful when a future
spreadsheet layout needs another adapter.
"""

from __future__ import annotations

import itertools
import warnings as python_warnings
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from time import perf_counter
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from lume_ingestion.errors import IngestionFailure
from lume_ingestion.models import SourceFile, Warning
from lume_ingestion.accounting import recognize_xlsx_document_type


_MAIN_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"


def _raw_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _merged_cells(path: Path) -> dict[str, list[str]]:
    """Read merge ranges from the package without abandoning read-only mode."""

    try:
        with zipfile.ZipFile(path) as archive:
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {
                relationship.attrib["Id"]: relationship.attrib["Target"].lstrip("/")
                for relationship in relationships.findall(f"{{{_PACKAGE_REL_NAMESPACE}}}Relationship")
            }
            result: dict[str, list[str]] = {}
            for sheet in workbook.findall(f".//{{{_MAIN_NAMESPACE}}}sheet"):
                relationship_id = sheet.attrib.get(f"{{{_REL_NAMESPACE}}}id")
                target = targets.get(relationship_id or "")
                if not target:
                    continue
                worksheet_path = target if target.startswith("xl/") else f"xl/{target}"
                worksheet = ElementTree.fromstring(archive.read(worksheet_path))
                result[sheet.attrib["name"]] = [
                    node.attrib["ref"]
                    for node in worksheet.findall(f".//{{{_MAIN_NAMESPACE}}}mergeCell")
                    if node.attrib.get("ref")
                ]
            return result
    except (KeyError, OSError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise IngestionFailure("invalid_xlsx", "A planilha XLSX nao pode ser lida como um pacote valido.", reason=str(exc)) from exc


def _used_columns(sheet: Any) -> int:
    """Ignore producer-added empty columns without losing populated headers."""

    return max((cell.column for row in sheet.iter_rows(max_row=min(sheet.max_row, 100)) for cell in row if cell.value is not None), default=1)


class XlsxParser:
    name = "xlsx-openpyxl"
    version = "0.1.0"

    def extract(
        self,
        path: Path,
        source: SourceFile,
        warnings: list[Warning],
        limits: Any = None,
    ) -> dict[str, Any]:
        del limits
        started = perf_counter()
        merged_cells = _merged_cells(path)
        try:
            # Some producer tools omit the default style.  It is harmless for
            # values and should not make a valid ledger noisy or fail.
            with python_warnings.catch_warnings():
                python_warnings.filterwarnings("ignore", message=r"Workbook contains no default style.*", category=UserWarning)
                formula_book = load_workbook(path, read_only=True, data_only=False)
                cached_book = load_workbook(path, read_only=True, data_only=True)
        except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
            raise IngestionFailure("invalid_xlsx", "A planilha XLSX esta corrompida ou possui estrutura invalida.", reason=str(exc)) from exc

        if formula_book.sheetnames != cached_book.sheetnames:
            raise IngestionFailure("invalid_xlsx", "As leituras de formulas e valores da planilha divergem nas abas.")

        formula_count = 0
        missing_formula_cache: list[str] = []
        sheets: list[dict[str, Any]] = []
        for formula_sheet, cached_sheet in zip(formula_book.worksheets, cached_book.worksheets, strict=True):
            rows: list[dict[str, Any]] = []
            max_column = 0
            max_row = 0
            max_column = _used_columns(formula_sheet)
            formula_rows = formula_sheet.iter_rows(max_col=max_column)
            cached_rows = cached_sheet.iter_rows(max_col=max_column)
            for guessed_row, (formula_cells, cached_cells) in enumerate(
                itertools.zip_longest(formula_rows, cached_rows, fillvalue=()), start=1
            ):
                if len(formula_cells) != len(cached_cells):
                    raise IngestionFailure(
                        "invalid_xlsx",
                        "As leituras de formulas e valores da planilha divergem nas celulas.",
                        sheet=formula_sheet.title,
                    )
                if not formula_cells:
                    continue
                row_number = next((int(cell.row) for cell in formula_cells if hasattr(cell, "row")), guessed_row)
                cells: list[dict[str, Any]] = []
                for column_index, (formula_cell, cached_cell) in enumerate(zip(formula_cells, cached_cells, strict=True), start=1):
                    column = int(getattr(formula_cell, "column", column_index))
                    coordinate = str(getattr(formula_cell, "coordinate", f"{get_column_letter(column)}{row_number}"))
                    is_formula = getattr(formula_cell, "data_type", None) == "f"
                    value = getattr(formula_cell, "value", None)
                    calculated_value = getattr(cached_cell, "value", None) if is_formula else None
                    if is_formula:
                        formula_count += 1
                        if calculated_value is None:
                            missing_formula_cache.append(coordinate)
                    if value is not None or is_formula:
                        cells.append(
                            {
                                "coordinate": coordinate,
                                "column": column,
                                "value": None if is_formula else _raw_value(value),
                                "formula": str(value) if is_formula else None,
                                "calculated_value": _raw_value(calculated_value),
                            }
                        )
                if cells:
                    max_row = max(max_row, row_number)
                    rows.append({"row_number": row_number, "cells": cells})
            sheets.append(
                {
                    "name": formula_sheet.title,
                    "dimensions": {"max_row": max_row, "max_column": max_column},
                    "merged_cells": merged_cells.get(formula_sheet.title, []),
                    "rows": rows,
                }
            )
        if missing_formula_cache:
            warnings.append(
                Warning(
                    code="formula_cache_missing",
                    message="Uma ou mais formulas nao possuem valor calculado em cache; elas permanecem no RAW.",
                    details={"cells": missing_formula_cache[:50], "count": len(missing_formula_cache)},
                )
            )
        if formula_count and (
            bool(getattr(formula_book.calculation, "fullCalcOnLoad", False))
            or bool(getattr(formula_book.calculation, "forceFullCalc", False))
            or getattr(formula_book.calculation, "calcMode", None) == "manual"
        ):
            warnings.append(
                Warning(
                    code="formula_cache_may_be_stale",
                    message="O arquivo solicita recalculo no Excel; os valores em cache podem estar desatualizados.",
                    details={"formula_count": formula_count},
                )
            )
        formula_book.close()
        cached_book.close()

        return {
            "schema_version": "1.0",
            "source": source.model_dump(mode="json"),
            "source_format": "xlsx",
            "document_type": recognize_xlsx_document_type(sheets) or "cash_ledger",
            "parser": self.name,
            "parser_version": self.version,
            "extraction_duration_ms": round((perf_counter() - started) * 1000),
            "workbook": {"sheet_count": len(sheets), "sheets": sheets, "formula_count": formula_count},
            "warnings": [warning.model_dump(mode="json") for warning in warnings],
            "errors": [],
        }
